import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation(r'd:\ebca\310_backup_original.pptx')
s19 = prs.slides[18]
print('Slide 19 Title:')
for shp in s19.shapes:
    if shp.has_text_frame and shp.text_frame.text.strip():
        print(' ', shp.text_frame.text.strip().replace('\n', ' ')[:60])

print('\nTotal shapes on Slide 19:', len(s19.shapes))
for i, shp in enumerate(s19.shapes):
    txt = shp.text_frame.text.replace('\n', ' ')[:30] if shp.has_text_frame else '<no text>'
    print(f'Shape {i}: type={shp.shape_type}, text="{txt}"')
