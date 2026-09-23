import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.dml.color import RGBColor

prs = Presentation(r'd:\ebca\310_backup_original.pptx')

for s_idx in [3, 4, 5, 6, 18]: # Slides 4, 5, 6, 7, 19
    slide = prs.slides[s_idx]
    title = ''
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip() and not title:
            title = shp.text_frame.text.strip().split('\n')[0]
    print(f"\n=== Slide {s_idx+1}: {title} (Total shapes: {len(slide.shapes)}) ===")
    color_counts = {}
    for shp in slide.shapes:
        if hasattr(shp, 'fill') and shp.fill.type is not None:
            try:
                clr = shp.fill.fore_color.rgb
                color_counts[str(clr)] = color_counts.get(str(clr), 0) + 1
            except Exception:
                pass
    print("Fill colors found:", color_counts)
