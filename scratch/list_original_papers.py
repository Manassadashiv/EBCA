import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'd:\ebca\310_backup_original.pptx')
for idx in range(7, 16):
    s = prs.slides[idx]
    title = ''
    for shp in s.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip() and not title:
            title = shp.text_frame.text.strip().split('\n')[0]
        if shp.has_table:
            print(f"\n{title} (Slide {idx+1}):")
            for r_idx, row in enumerate(shp.table.rows):
                if r_idx == 0: continue
                author = row.cells[0].text.replace('\n', ' ').strip()
                title_paper = row.cells[1].text.replace('\n', ' ').strip()
                print(f"  {author} | {title_paper}")
