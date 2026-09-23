import pptx, re

prs = pptx.Presentation('d:/ebca/scratch/test_fixed_numbers.pptx')
with open('d:/ebca/scratch/all_stage_patterns.txt', 'w', encoding='utf-8') as out:
    for idx, s in enumerate(prs.slides):
        slide_num = idx + 1
        for sh in s.shapes:
            txt = ''
            if sh.has_text_frame:
                txt = sh.text_frame.text
            elif sh.has_table:
                txt = ' '.join(c.text_frame.text for row in sh.table.rows for c in row.cells)
            m = re.findall(r'Stage \d{2}[A-Za-z]?', txt)
            if m:
                out.write(f'Slide {slide_num}: {m} in "{txt[:100]}"\n')
print("Done checking all stage patterns!")
