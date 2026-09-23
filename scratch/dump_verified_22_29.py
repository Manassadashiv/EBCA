import pptx

prs = pptx.Presentation('D:/ebca/310_Stage3_Updated_Final.pptx')
with open('d:/ebca/scratch/verified_slides_22_29.txt', 'w', encoding='utf-8') as out:
    for idx in [21, 22, 23, 25, 26, 27, 28]: # Slides 22, 23, 24, 26, 27, 28, 29
        s = prs.slides[idx]
        out.write(f"=== SLIDE {idx+1} ===\n")
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        out.write(f"  [text] {t}\n")
            elif sh.has_table:
                for r_idx, row in enumerate(sh.table.rows):
                    row_txt = " | ".join(c.text_frame.text.strip() for c in row.cells)
                    out.write(f"  [table row {r_idx}] {row_txt}\n")
print("Dumped verified slides 22-29!")
