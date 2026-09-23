import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'd:\ebca\310_backup_original.pptx')
slide = prs.slides[3] # Slide 4

line_colors = {}
text_colors = {}
for shp in slide.shapes:
    if hasattr(shp, 'line') and shp.line is not None:
        try:
            c = str(shp.line.color.rgb)
            line_colors[c] = line_colors.get(c, 0) + 1
        except Exception:
            pass
    if shp.has_text_frame:
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                try:
                    c = str(r.font.color.rgb)
                    text_colors[c] = text_colors.get(c, 0) + 1
                except Exception:
                    pass

print("Line colors on slide 4:", line_colors)
print("Text colors on slide 4:", text_colors)
