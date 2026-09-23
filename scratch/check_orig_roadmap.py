import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Let's inspect old slide 19 from 310_backup_original.pptx to copy its exact text into our Slide 14 builder!
prs_orig = Presentation(r"d:\ebca\310_backup_original.pptx")
s19_orig = prs_orig.slides[18]
print("Roadmap title:", s19_orig.shapes[0].text if s19_orig.shapes[0].has_text_frame else "None")
for i, shp in enumerate(s19_orig.shapes):
    if shp.has_text_frame and shp.text_frame.text.strip():
        txt = shp.text_frame.text.replace("\n", " -- ")[:80]
        print(f"Shape {i}: {txt}")
