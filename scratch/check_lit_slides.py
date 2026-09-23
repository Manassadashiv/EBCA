import pptx, re

prs = pptx.Presentation('d:/ebca/scratch/test_fixed_numbers.pptx')
with open('d:/ebca/scratch/lit_stage_matches.txt', 'w', encoding='utf-8') as out:
    for idx in range(2, 11):
        s = prs.slides[idx]
        for sh in s.shapes:
            if sh.has_table:
                for r_idx, r in enumerate(sh.table.rows):
                    for c_idx, c in enumerate(r.cells):
                        txt = c.text_frame.text
                        m = re.findall(r'Stage \d{2}[A-Za-z]?', txt)
                        if m:
                            out.write(f'Slide {idx+1} cell ({r_idx},{c_idx}): {m} in "{txt}"\n')
print("Done checking lit slides!")
