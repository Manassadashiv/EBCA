import sys
from pptx import Presentation

prs = Presentation(r'd:\ebca\test_bw_output.pptx')
print(f"Initial slides count: {len(prs.slides)}")

def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slide_id = xml_slides[index]
    r_id = slide_id.rId
    prs.part.drop_rel(r_id)
    del xml_slides[index]

# Delete last slide
delete_slide(prs, len(prs.slides) - 1)
print(f"After deleting last slide: {len(prs.slides)}")
prs.save(r'd:\ebca\test_del_output.pptx')

prs2 = Presentation(r'd:\ebca\test_del_output.pptx')
print(f"Reloaded slides count: {len(prs2.slides)}")
