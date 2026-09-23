import pptx

prs = pptx.Presentation('D:/ebca/310_Stage3_Updated_Final.pptx')
s28 = prs.slides[27]
with open('d:/ebca/scratch/s28_runs_output.txt', 'w', encoding='utf-8') as out:
    for sh in s28.shapes:
        if sh.has_table:
            for r_i, row in enumerate(sh.table.rows):
                for c_i, cell in enumerate(row.cells):
                    runs_info = [(r.text, r.font.size.pt if r.font.size else None) for p in cell.text_frame.paragraphs for r in p.runs]
                    out.write(f'Cell ({r_i},{c_i}): full="{cell.text_frame.text}" runs={runs_info}\n')
print("Dumped s28 table runs!")
