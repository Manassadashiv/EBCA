import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'd:\ebca\310_backup_original.pptx')
for idx in [19, 20, 21]:
    s = prs.slides[idx]
    print(f'=== Slide {idx+1} ===')
    for i, shp in enumerate(s.shapes):
        txt = shp.text_frame.text.replace('\n', ' ')[:80] if shp.has_text_frame else '<no text>'
        print(f'  Shape {i}: type={shp.shape_type}, text="{txt}"')
