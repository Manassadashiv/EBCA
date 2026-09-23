import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r"C:\Users\MANAS\OneDrive\Desktop\ppt\19_08\310.pptx")
print(f"Total slides in generated presentation: {len(prs.slides)}\n")

for idx, s in enumerate(prs.slides):
    title = ''
    pics = 0
    tables = 0
    for shp in s.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip() and not title:
            lines = [l.strip() for l in shp.text_frame.text.split('\n') if l.strip()]
            if lines:
                title = lines[0]
        if shp.shape_type == 13: # Picture
            pics += 1
        if shp.has_table:
            tables += 1
    print(f"Slide {idx+1:02d}: {title[:55]:<55} | Pics: {pics} | Tables: {tables} | Shapes: {len(s.shapes)}")
