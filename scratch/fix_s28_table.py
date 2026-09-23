import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = pptx.Presentation('d:/ebca/scratch/test_fixed_numbers.pptx')
s28 = prs.slides[27]

gate_data = [
    ("Gate 1: Body Proprioception", "Forward Velocity (v) R²", "≥ 0.90", "R² = 0.9613", "PASS"),
    ("Gate 2: Angular Gyroscope", "Rotational Velocity (ω) R²", "≥ 0.80", "R² = 0.8466", "PASS"),
    ("Gate 3: Dynamic Sonar Range", "Virtual Rays Range R²", "≥ 0.50", "R² = 0.5377", "PASS"),
    ("Gate 4: Collision Prediction", "Imminent Contact ROC-AUC", "≥ 0.90", "ROC-AUC = 0.995", "PASS"),
    ("Gate 5: Multi-Step Stability", "5-Step Rollout Drift", "< 15.0%", "Drift = 8.18%", "PASS"),
    ("Gate 6: Causal Sensitivity", "Action Intervention Δ", "> 0.35", "Δ = 0.442", "PASS"),
    ("Gate 7: Compute Latency", "Inference Cycle Time", "≤ 10.0 ms", "2.12 ms", "PASS"),
]

for sh in s28.shapes:
    if sh.has_table and len(sh.table.rows) == 8:
        tbl = sh.table
        for r_idx, row in enumerate(gate_data):
            for c_idx, val in enumerate(row):
                cell = tbl.cell(r_idx + 1, c_idx)
                tf = cell.text_frame
                p = tf.paragraphs[0]
                p.text = ""
                r = p.add_run()
                r.text = val
                r.font.name = "Times New Roman"
                r.font.size = Pt(18)
                r.font.bold = (c_idx == 4 or "R²" in val or "ROC" in val)
                r.font.color.rgb = RGBColor(0, 0, 0)
                if c_idx in [0, 4] or len(val) <= 12:
                    p.alignment = PP_ALIGN.CENTER
                else:
                    p.alignment = PP_ALIGN.LEFT

prs.save('d:/ebca/scratch/test_fixed_numbers.pptx')
print("Slide 28 table cells perfectly reset!")
