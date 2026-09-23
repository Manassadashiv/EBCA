import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation(r'd:\ebca\310_backup_original.pptx')
print("Slide count:", len(prs.slides))

# Check table formatting in slide 8
s8 = prs.slides[7]
for shp in s8.shapes:
    if shp.has_table:
        tbl = shp.table
        print(f"Table rows: {len(tbl.rows)}, cols: {len(tbl.columns)}")
        print("Col widths:", [c.width.inches for c in tbl.columns])
        for r_idx, row in enumerate(tbl.rows):
            print(f"Row {r_idx}: {[c.text.replace(chr(10), ' ')[:25] for c in row.cells]}")
