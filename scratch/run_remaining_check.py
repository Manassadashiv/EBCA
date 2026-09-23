import pptx

prs = pptx.Presentation('d:/ebca/scratch/test_fixed_numbers.pptx')
with open('d:/ebca/scratch/remaining_check.txt', 'w', encoding='utf-8') as out:
    for idx, s in enumerate(prs.slides):
        slide_num = idx + 1
        for sh in s.shapes:
            txt = ''
            if sh.has_text_frame:
                txt = sh.text_frame.text
            elif sh.has_table:
                txt = ' '.join(c.text_frame.text for row in sh.table.rows for c in row.cells)
            for bad in ['00650', '09613', '08466', '05377', '0995', '818%', '150%', '0442', '212 ms', '100 ms', '022 m', 'Stage 21', 'Stage 22', '035 where', '005, 10', '005-025', '005–025', '982%', '050-10', '050–10', '241%', '01-05', '01–05', '02-15', '02–15', '0065 m', '0071 m', '0062 m', '0059 m', '0057 m', '002 rad/s']:
                if bad in txt:
                    out.write(f'Slide {slide_num} matched "{bad}": {txt[:80]}\n')
print("Done writing remaining_check.txt!")
