r"""
tools/generate_slide19_roadmap.py — Generates Slide 19: Developmental Phase Roadmap

Produces:
1. CARL_SLIDE_19_DEVELOPMENT_ROADMAP.png (High-Resolution 16:9 Presentation Image)
2. CARL_SLIDE_19_DEVELOPMENT_ROADMAP.pptx (Editable 16:9 PowerPoint Slide)
3. CARL_SLIDE_19_ROADMAP_NOTES.md (Explanatory Architectural Reference Notes)
"""

import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as patches

REPO_ROOT = r"D:\ebca"

# -----------------------------------------------------------------------------
# 1. GENERATE HIGH-RESOLUTION PNG VIA MATPLOTLIB
# -----------------------------------------------------------------------------
def generate_roadmap_png():
    print(">>> Rendering CARL_SLIDE_19_DEVELOPMENT_ROADMAP.png (300 DPI, 16:9)...")
    
    # 16:9 aspect ratio at 16 inches x 9 inches
    fig, ax = plt.subplots(figsize=(16, 9), dpi=300)
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 9)
    ax.axis('off')
    
    # White background
    fig.patch.set_facecolor('#ffffff')
    ax.set_facecolor('#ffffff')
    
    # -------------------------------------------------------------------------
    # Top Header & Legend
    # -------------------------------------------------------------------------
    # Category Tag
    ax.text(0.6, 8.45, "CARL-DX ARCHITECTURAL METHODOLOGY", fontsize=9, fontweight='bold', color='#0284c7', family='sans-serif')
    
    # Main Title
    ax.text(0.6, 8.10, "CARL DEVELOPMENT & IMPLEMENTATION ROADMAP", fontsize=18, fontweight='heavy', color='#0f172a', family='sans-serif')
    
    # Subtitle
    ax.text(0.6, 7.78, "From the full CARL research vision to a validated embodied foundation and subsequent developmental phases", fontsize=10.5, color='#475569', family='sans-serif')
    
    # Status Legend (Top Right)
    legend_bg = patches.FancyBboxPatch((10.6, 7.72), 4.8, 0.85, boxstyle="round,pad=0.08,rounding_size=0.08", facecolor='#f8fafc', edgecolor='#cbd5e1', linewidth=1.0)
    ax.add_patch(legend_bg)
    
    ax.text(10.75, 8.35, "STATUS LEGEND:", fontsize=8, fontweight='bold', color='#334155', family='sans-serif')
    
    # Legend item 1: Implemented
    p0_dot = patches.Circle((10.85, 8.0), 0.09, facecolor='#16a34a', edgecolor='none')
    ax.add_patch(p0_dot)
    ax.text(11.05, 7.93, "COMPLETE (P0 Validated)", fontsize=7.5, fontweight='bold', color='#166534', family='sans-serif')
    
    # Legend item 2: Next Phase
    p1_dot = patches.Circle((12.75, 8.0), 0.09, facecolor='#0284c7', edgecolor='none')
    ax.add_patch(p1_dot)
    ax.text(12.95, 7.93, "NEXT / UNLOCKED (P1)", fontsize=7.5, fontweight='bold', color='#0369a1', family='sans-serif')
    
    # Legend item 3: Future Roadmap
    pf_dot = patches.Circle((14.50, 8.0), 0.09, facecolor='#94a3b8', edgecolor='none')
    ax.add_patch(pf_dot)
    ax.text(14.70, 7.93, "ROADMAP (P2-P8)", fontsize=7.5, fontweight='bold', color='#475569', family='sans-serif')
    
    # -------------------------------------------------------------------------
    # Phase Data Definitions
    # -------------------------------------------------------------------------
    phases = [
        # Tier 1 (0 to 4)
        {
            "id": "P0",
            "name": "EMBODIED FOUNDATION",
            "status": "COMPLETE [VALIDATED]",
            "status_type": "COMPLETE",
            "bullets": [
                "• 6-DoF MuJoCo physical embodiment",
                "• Active Inference & Free Energy G(a)",
                "• Proprioception + 15-ray Sonar array",
                "• Interoceptive homeostasis (Energy/Heat)",
                "• HDC episodic memory & online friction",
                "• Safety veto & 81/81 regression certified"
            ],
            "result": "Result: Validated P0 Baseline",
            "col": 0, "row": 0
        },
        {
            "id": "P1",
            "name": "INDIVIDUAL EMERGENCE",
            "status": "NEXT / UNLOCKED",
            "status_type": "NEXT",
            "bullets": [
                "• Experience -> error -> adaptation",
                "• Memory-driven behavioral change",
                "• Controlled emergence laboratory",
                "• Progressive biological integration:",
                "   - LSM Liquid State Reservoir",
                "   - STDP continual synaptic plasticity",
                "   - 4-factor endocrine modulation"
            ],
            "result": "Focus: Developmental Learning",
            "col": 1, "row": 0
        },
        {
            "id": "P2",
            "name": "SENSORY & SPATIAL",
            "status": "FUTURE ROADMAP",
            "status_type": "FUTURE",
            "bullets": [
                "• Latent visual world model",
                "• Active optical gaze scanning",
                "• Spatial cognitive representation",
                "• Hippocampal navigation core",
                "• Place-cell & grid-cell firing arrays"
            ],
            "result": "Target: Latent Spatial Cognition",
            "col": 2, "row": 0
        },
        {
            "id": "P3",
            "name": "ADAPTIVE ROUTING",
            "status": "FUTURE ROADMAP",
            "status_type": "FUTURE",
            "bullets": [
                "• Continuous trajectory optimization",
                "• Long-horizon path planning",
                "• Multi-step counterfactual pruning",
                "• Physarum route reinforcement",
                "• Tubular flow network adaptation"
            ],
            "result": "Target: Complex Maze Mastery",
            "col": 3, "row": 0
        },
        {
            "id": "P4",
            "name": "SOCIAL CARL",
            "status": "FUTURE ROADMAP",
            "status_type": "FUTURE",
            "bullets": [
                "• Multi-agent co-existence arena",
                "• Peer observation & tracking",
                "• Non-verbal acoustic/visual signals",
                "• Mirror-neuron-inspired imitation",
                "• Inter-agent behavioral resonance"
            ],
            "result": "Target: Peer Social Learning",
            "col": 4, "row": 0
        },
        
        # Tier 2 (5 to 8)
        {
            "id": "P5",
            "name": "COLLECTIVE SWARM",
            "status": "FUTURE ROADMAP",
            "status_type": "FUTURE",
            "bullets": [
                "• Selective shared HDC memory sync",
                "• Swarm hazard & knowledge broadcast",
                "• Collective route & resource discovery",
                "• Emergent division of labor",
                "• Cooperative multi-organism goals"
            ],
            "result": "Target: Collective Intelligence",
            "col": 0, "row": 1
        },
        {
            "id": "P6",
            "name": "METACOGNITION",
            "status": "FUTURE ROADMAP",
            "status_type": "FUTURE",
            "bullets": [
                "• Failure-pattern self-monitoring",
                "• Higher-order counterfactual audit",
                "• Causal attribution & self-correction",
                "• Epistemic confidence estimation",
                "• Goal arbitration & intrinsic priors"
            ],
            "result": "Target: Causal Self-Reasoning",
            "col": 1, "row": 1
        },
        {
            "id": "P7",
            "name": "MORPHOLOGICAL CO-EVO",
            "status": "FUTURE ROADMAP",
            "status_type": "FUTURE",
            "bullets": [
                "• Brain-directed morphology adaptation",
                "• Body-brain developmental co-tuning",
                "• Dynamic sensor placement evolution",
                "• Actuator torque/geometry scaling",
                "• Evolutionary morphology gauntlets"
            ],
            "result": "Target: Embodied Co-Evolution",
            "col": 2, "row": 1
        },
        {
            "id": "P8",
            "name": "PHYSICAL EMBODIMENT",
            "status": "FUTURE ROADMAP",
            "status_type": "FUTURE",
            "bullets": [
                "• Sim-to-Real parameter transfer",
                "• Physical brushless DC actuation",
                "• Embedded STM32/Jetson deployment",
                "• Real-world acoustic & optical sonar",
                "• Physical companion validation"
            ],
            "result": "Target: Physical Reality Transfer",
            "col": 3, "row": 1
        }
    ]
    
    # -------------------------------------------------------------------------
    # Layout Coordinates
    # -------------------------------------------------------------------------
    box_w = 2.80
    box_h = 2.45
    col_x = [0.60, 3.65, 6.70, 9.75, 12.80]
    row_y = [4.80, 2.05]
    
    for p in phases:
        c_idx = p["col"]
        r_idx = p["row"]
        bx = col_x[c_idx]
        by = row_y[r_idx]
        
        # Color Styling based on status
        if p["status_type"] == "COMPLETE":
            card_bg = '#f0fdf4' # Light green
            card_border = '#16a34a' # Dark green
            badge_bg = '#16a34a'
            badge_fg = '#ffffff'
            title_col = '#14532d'
            lw = 2.2
        elif p["status_type"] == "NEXT":
            card_bg = '#f0f9ff' # Light blue
            card_border = '#0284c7' # Vivid blue
            badge_bg = '#0284c7'
            badge_fg = '#ffffff'
            title_col = '#0c4a6e'
            lw = 2.2
        else:
            card_bg = '#ffffff' # Clean white
            card_border = '#cbd5e1' # Slate border
            badge_bg = '#e2e8f0'
            badge_fg = '#475569'
            title_col = '#1e293b'
            lw = 1.2
            
        # Draw Card Box
        card_patch = patches.FancyBboxPatch((bx, by), box_w, box_h, boxstyle="round,pad=0.08,rounding_size=0.10", facecolor=card_bg, edgecolor=card_border, linewidth=lw, zorder=2)
        ax.add_patch(card_patch)
        
        # Phase Header: ID + Name
        ax.text(bx + 0.12, by + box_h - 0.28, f"{p['id']} — {p['name']}", fontsize=9.5, fontweight='bold', color=title_col, family='sans-serif', zorder=3)
        
        # Status Badge
        badge_w = 1.30 if "NEXT" in p["status"] else (1.10 if "COMPLETE" in p["status"] else 1.20)
        badge_patch = patches.FancyBboxPatch((bx + box_w - badge_w - 0.08, by + box_h - 0.35), badge_w, 0.24, boxstyle="round,pad=0.04,rounding_size=0.06", facecolor=badge_bg, edgecolor='none', zorder=3)
        ax.add_patch(badge_patch)
        ax.text(bx + box_w - (badge_w/2.0) - 0.08, by + box_h - 0.25, p["status"], fontsize=6.8, fontweight='bold', color=badge_fg, ha='center', va='center', family='sans-serif', zorder=4)
        
        # Separator line
        ax.plot([bx + 0.10, bx + box_w - 0.10], [by + box_h - 0.42, by + box_h - 0.42], color=card_border, linewidth=0.8, zorder=3)
        
        # Bullet Points
        text_y = by + box_h - 0.62
        for b in p["bullets"]:
            b_font_w = 'bold' if b.startswith("• Progressive") else 'normal'
            b_col = '#0f172a' if b.startswith("• Progressive") else '#334155'
            ax.text(bx + 0.12, text_y, b, fontsize=7.2, fontweight=b_font_w, color=b_col, family='sans-serif', zorder=3)
            text_y -= 0.25
            
        # Result Bottom Line
        res_bg = patches.FancyBboxPatch((bx + 0.08, by + 0.08), box_w - 0.16, 0.24, boxstyle="round,pad=0.02,rounding_size=0.04", facecolor='#f8fafc', edgecolor='#e2e8f0', linewidth=0.8, zorder=3)
        ax.add_patch(res_bg)
        res_col = '#15803d' if p["status_type"] == "COMPLETE" else ('#0369a1' if p["status_type"] == "NEXT" else '#475569')
        ax.text(bx + 0.14, by + 0.18, p["result"], fontsize=7.0, fontweight='bold', color=res_col, family='sans-serif', zorder=4)
        
    # -------------------------------------------------------------------------
    # Developmental Flow Connectors (Arrows)
    # -------------------------------------------------------------------------
    # Tier 1 Horizontal Arrows: P0 -> P1 -> P2 -> P3 -> P4
    for i in range(4):
        ax.annotate("", xy=(col_x[i+1] - 0.04, row_y[0] + box_h/2), xytext=(col_x[i] + box_w + 0.04, row_y[0] + box_h/2),
                    arrowprops=dict(arrowstyle="-|>", color='#0284c7' if i == 0 else '#94a3b8', lw=2.2 if i == 0 else 1.6, mutation_scale=14), zorder=5)
                    
    # Tier 1 to Tier 2 Transition Arrow (P4 -> P5 wrap around)
    # Arrow from P4 down-left to P5
    p4_bottom_x = col_x[4] + box_w/2
    p4_bottom_y = row_y[0]
    p5_top_x = col_x[0] + box_w/2
    p5_top_y = row_y[1] + box_h
    
    # Draw smooth transition curve from P4 to P5
    ax.annotate("", xy=(col_x[0] + box_w/2, row_y[1] + box_h + 0.05), xytext=(col_x[4] + box_w/2, row_y[0] - 0.05),
                arrowprops=dict(arrowstyle="-|>", connectionstyle="arc3,rad=-0.15", color='#94a3b8', lw=1.6, linestyle='--', mutation_scale=14), zorder=5)
    ax.text(8.2, 4.62, "Developmental Continuum (P4 → P5)", fontsize=7.2, fontstyle='italic', color='#64748b', ha='center', backgroundcolor='#ffffff')

    # Tier 2 Horizontal Arrows: P5 -> P6 -> P7 -> P8
    for i in range(3):
        ax.annotate("", xy=(col_x[i+1] - 0.04, row_y[1] + box_h/2), xytext=(col_x[i] + box_w + 0.04, row_y[1] + box_h/2),
                    arrowprops=dict(arrowstyle="-|>", color='#94a3b8', lw=1.6, mutation_scale=14), zorder=5)
                    
    # Tier 2 Right Box: Developmental Paradigm & Governance Box
    gov_x = col_x[4]
    gov_y = row_y[1]
    gov_patch = patches.FancyBboxPatch((gov_x, gov_y), box_w, box_h, boxstyle="round,pad=0.08,rounding_size=0.10", facecolor='#f8fafc', edgecolor='#0284c7', linewidth=1.5, linestyle=':', zorder=2)
    ax.add_patch(gov_patch)
    
    ax.text(gov_x + 0.14, gov_y + box_h - 0.28, "RESEARCH VISION GOVERNANCE", fontsize=8.8, fontweight='bold', color='#0f172a', family='sans-serif', zorder=3)
    ax.plot([gov_x + 0.10, gov_x + box_w - 0.10], [gov_y + box_h - 0.40, gov_y + box_h - 0.40], color='#cbd5e1', linewidth=0.8, zorder=3)
    
    gov_bullets = [
        "• Full CARL Vision is Progressive:",
        "  Not all modules exist simultaneously.",
        "• Strict Causal Ablation Protocol:",
        "  Every new mechanism must be",
        "  quantitatively validated vs baseline.",
        "• Zero Scripted Competence:",
        "  Only emergent, self-organized",
        "  adaptation is scientifically valid."
    ]
    gy = gov_y + box_h - 0.62
    for gb in gov_bullets:
        ax.text(gov_x + 0.12, gy, gb, fontsize=7.2, color='#334155', family='sans-serif', zorder=3)
        gy -= 0.24
        
    # -------------------------------------------------------------------------
    # Bottom Footer Banner: Objective Mapping to Slide 3
    # -------------------------------------------------------------------------
    footer_box = patches.FancyBboxPatch((0.60, 0.40), 15.00, 1.40, boxstyle="round,pad=0.06,rounding_size=0.08", facecolor='#f1f5f9', edgecolor='#94a3b8', linewidth=1.2, zorder=2)
    ax.add_patch(footer_box)
    
    ax.text(0.80, 1.52, "ORIGINAL 8 RESEARCH OBJECTIVES MAPPING (ALIGNMENT TO SLIDE 3)", fontsize=9.0, fontweight='heavy', color='#0f172a', family='sans-serif', zorder=3)
    ax.text(10.8, 1.52, "Principle: Baseline qualified in P0; Biological complexity unlocked in P1+", fontsize=7.5, fontstyle='italic', color='#475569', family='sans-serif', zorder=3)
    
    # 8 Objective Pills (2 rows of 4 pills)
    obj_data = [
        ("Obj 1: Integrated CARL / MuJoCo 6-DoF", "-> P0 VALIDATED", '#16a34a', '#14532d', '#f0fdf4'),
        ("Obj 2: LSM Reservoir + 4-Factor Endocrine", "-> P1 (Next / Progressive)", '#0284c7', '#0c4a6e', '#f0f9ff'),
        ("Obj 3: STDP Continual Synaptic Plasticity", "-> P1 (Next / Progressive)", '#0284c7', '#0c4a6e', '#f0f9ff'),
        ("Obj 4: Hippocampal Grid & Place Cells", "-> P2 (Sensory & Spatial)", '#64748b', '#334155', '#ffffff'),
        ("Obj 5: Free-Energy Active Inference", "-> P0 VALIDATED", '#16a34a', '#14532d', '#f0fdf4'),
        ("Obj 6: Physarum Route Reinforcement", "-> P3 (Adaptive Routing)", '#64748b', '#334155', '#ffffff'),
        ("Obj 7: Mirror-Neuron / Social Imitation", "-> P4 (Social CARL)", '#64748b', '#334155', '#ffffff'),
        ("Obj 8: 30 Hz Digital-Twin Validation", "-> P0 Baseline (P1+ Ongoing)", '#16a34a', '#14532d', '#f0fdf4')
    ]
    
    pw = 3.55
    ph = 0.38
    px_coords = [0.80, 4.45, 8.10, 11.75]
    py_coords = [0.98, 0.50]
    
    for idx, (label, target, border_c, text_c, bg_c) in enumerate(obj_data):
        col_i = idx % 4
        row_i = idx // 4
        px = px_coords[col_i]
        py = py_coords[row_i]
        
        # Pill patch
        pill = patches.FancyBboxPatch((px, py), pw, ph, boxstyle="round,pad=0.03,rounding_size=0.06", facecolor=bg_c, edgecolor=border_c, linewidth=1.0, zorder=3)
        ax.add_patch(pill)
        
        ax.text(px + 0.10, py + 0.20, label, fontsize=6.8, fontweight='bold', color='#1e293b', family='sans-serif', zorder=4)
        ax.text(px + pw - 0.10, py + 0.20, target, fontsize=6.8, fontweight='bold', color=text_c, ha='right', family='sans-serif', zorder=4)
        
    out_png_path = os.path.join(REPO_ROOT, "CARL_SLIDE_19_DEVELOPMENT_ROADMAP.png")
    plt.savefig(out_png_path, bbox_inches='tight', pad_inches=0.1)
    plt.close()
    print(f"  [OK] Saved High-Resolution PNG -> {out_png_path}")
    return out_png_path

# -----------------------------------------------------------------------------
# 2. GENERATE EDITABLE PPTX SLIDE VIA PYTHON-PPTX
# -----------------------------------------------------------------------------
def generate_roadmap_pptx():
    print(">>> Rendering CARL_SLIDE_19_DEVELOPMENT_ROADMAP.pptx (16:9 Editable PPTX)...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.500)
    
    # Use blank slide layout (index 6 in default template)
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Background: Solid White
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)
    
    # -------------------------------------------------------------------------
    # Header & Legend
    # -------------------------------------------------------------------------
    # Title Text Box
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8.5), Inches(1.1))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p_tag = tf.paragraphs[0]
    p_tag.text = "CARL-DX ARCHITECTURAL METHODOLOGY"
    p_tag.font.size = Pt(8)
    p_tag.font.bold = True
    p_tag.font.color.rgb = RGBColor(2, 132, 199) # Vivid Blue
    
    p_title = tf.add_paragraph()
    p_title.text = "CARL DEVELOPMENT & IMPLEMENTATION ROADMAP"
    p_title.font.size = Pt(18)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(15, 23, 42) # Slate-900
    
    p_sub = tf.add_paragraph()
    p_sub.text = "From the full CARL research vision to a validated embodied foundation and subsequent developmental phases"
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = RGBColor(71, 85, 105) # Slate-600
    
    # Legend Box (Top Right)
    leg_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(0.45), Inches(3.6), Inches(0.85))
    leg_box.fill.solid()
    leg_box.fill.fore_color.rgb = RGBColor(248, 250, 252)
    leg_box.line.color.rgb = RGBColor(203, 213, 225)
    leg_box.line.width = Pt(1)
    
    ltf = leg_box.text_frame
    ltf.word_wrap = True
    ltf.margin_left = ltf.margin_top = ltf.margin_right = ltf.margin_bottom = Inches(0.08)
    lp0 = ltf.paragraphs[0]
    lp0.text = "STATUS LEGEND:"
    lp0.font.size = Pt(8)
    lp0.font.bold = True
    lp0.font.color.rgb = RGBColor(51, 65, 85)
    
    lp1 = ltf.add_paragraph()
    lp1.text = "● COMPLETE (P0 Validated)   ● NEXT / UNLOCKED (P1)   ● ROADMAP (P2-P8)"
    lp1.font.size = Pt(7.5)
    lp1.font.bold = True
    lp1.font.color.rgb = RGBColor(71, 85, 105)
    
    # -------------------------------------------------------------------------
    # Phase Boxes Setup
    # -------------------------------------------------------------------------
    # We create 9 phase cards across 2 rows
    # Row 1: P0, P1, P2, P3, P4
    # Row 2: P5, P6, P7, P8, Governance
    
    phases = [
        # Row 1
        {
            "id": "P0", "name": "EMBODIED FOUNDATION", "status": "COMPLETE ✅", "type": "COMPLETE",
            "bullets": [
                "• 6-DoF MuJoCo physical embodiment",
                "• Active Inference & Free Energy G(a)",
                "• Proprioception + 15-ray Sonar array",
                "• Interoceptive homeostasis (Energy/Heat)",
                "• HDC episodic memory & online friction",
                "• Safety veto & 81/81 regression certified"
            ],
            "res": "Result: Validated P0 Baseline", "col": 0, "row": 0
        },
        {
            "id": "P1", "name": "INDIVIDUAL EMERGENCE", "status": "NEXT / UNLOCKED 🔓", "type": "NEXT",
            "bullets": [
                "• Experience → error → adaptation",
                "• Memory-driven behavioral change",
                "• Controlled emergence laboratory",
                "• Progressive biological integration:",
                "   - LSM Liquid State Reservoir",
                "   - STDP continual synaptic plasticity",
                "   - 4-factor endocrine modulation"
            ],
            "res": "Focus: Developmental Learning", "col": 1, "row": 0
        },
        {
            "id": "P2", "name": "SENSORY & SPATIAL", "status": "FUTURE ROADMAP", "type": "FUTURE",
            "bullets": [
                "• Latent visual world model",
                "• Active optical gaze scanning",
                "• Spatial cognitive representation",
                "• Hippocampal navigation core",
                "• Place-cell & grid-cell firing arrays"
            ],
            "res": "Target: Latent Spatial Cognition", "col": 2, "row": 0
        },
        {
            "id": "P3", "name": "ADAPTIVE ROUTING", "status": "FUTURE ROADMAP", "type": "FUTURE",
            "bullets": [
                "• Continuous trajectory optimization",
                "• Long-horizon path planning",
                "• Multi-step counterfactual pruning",
                "• Physarum route reinforcement",
                "• Tubular flow network adaptation"
            ],
            "res": "Target: Complex Maze Mastery", "col": 3, "row": 0
        },
        {
            "id": "P4", "name": "SOCIAL CARL", "status": "FUTURE ROADMAP", "type": "FUTURE",
            "bullets": [
                "• Multi-agent co-existence arena",
                "• Peer observation & tracking",
                "• Non-verbal acoustic/visual signals",
                "• Mirror-neuron-inspired imitation",
                "• Inter-agent behavioral resonance"
            ],
            "res": "Target: Peer Social Learning", "col": 4, "row": 0
        },
        # Row 2
        {
            "id": "P5", "name": "COLLECTIVE SWARM", "status": "FUTURE ROADMAP", "type": "FUTURE",
            "bullets": [
                "• Selective shared HDC memory sync",
                "• Swarm hazard & knowledge broadcast",
                "• Collective route & resource discovery",
                "• Emergent division of labor",
                "• Cooperative multi-organism goals"
            ],
            "res": "Target: Collective Intelligence", "col": 0, "row": 1
        },
        {
            "id": "P6", "name": "METACOGNITION", "status": "FUTURE ROADMAP", "type": "FUTURE",
            "bullets": [
                "• Failure-pattern self-monitoring",
                "• Higher-order counterfactual audit",
                "• Causal attribution & self-correction",
                "• Epistemic confidence estimation",
                "• Goal arbitration & intrinsic priors"
            ],
            "res": "Target: Causal Self-Reasoning", "col": 1, "row": 1
        },
        {
            "id": "P7", "name": "MORPHOLOGICAL CO-EVO", "status": "FUTURE ROADMAP", "type": "FUTURE",
            "bullets": [
                "• Brain-directed morphology adaptation",
                "• Body-brain developmental co-tuning",
                "• Dynamic sensor placement evolution",
                "• Actuator torque/geometry scaling",
                "• Evolutionary morphology gauntlets"
            ],
            "res": "Target: Embodied Co-Evolution", "col": 2, "row": 1
        },
        {
            "id": "P8", "name": "PHYSICAL EMBODIMENT", "status": "FUTURE ROADMAP", "type": "FUTURE",
            "bullets": [
                "• Sim-to-Real parameter transfer",
                "• Physical brushless DC actuation",
                "• Embedded STM32/Jetson deployment",
                "• Real-world acoustic & optical sonar",
                "• Physical companion validation"
            ],
            "res": "Target: Physical Reality Transfer", "col": 3, "row": 1
        },
        {
            "id": "GOV", "name": "VISION GOVERNANCE", "status": "SCIENTIFIC RIGOR", "type": "GOV",
            "bullets": [
                "• Full CARL Vision is Progressive:",
                "  Not all modules exist simultaneously.",
                "• Strict Causal Ablation Protocol:",
                "  Every mechanism tested vs baseline.",
                "• Zero Scripted Competence:",
                "  Only emergent adaptation is valid."
            ],
            "res": "Standard: Falsifiable Science", "col": 4, "row": 1
        }
    ]
    
    col_x_in = [0.50, 3.05, 5.60, 8.15, 10.70]
    row_y_in = [1.50, 3.75]
    card_w_in = 2.35
    card_h_in = 2.05
    
    for p in phases:
        cx = col_x_in[p["col"]]
        cy = row_y_in[p["row"]]
        
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(cy), Inches(card_w_in), Inches(card_h_in))
        card.fill.solid()
        
        if p["type"] == "COMPLETE":
            card.fill.fore_color.rgb = RGBColor(240, 253, 244) # Green tint
            card.line.color.rgb = RGBColor(22, 163, 74) # Green border
            card.line.width = Pt(2)
            title_rgb = RGBColor(20, 83, 45)
        elif p["type"] == "NEXT":
            card.fill.fore_color.rgb = RGBColor(240, 249, 255) # Blue tint
            card.line.color.rgb = RGBColor(2, 132, 199) # Blue border
            card.line.width = Pt(2)
            title_rgb = RGBColor(12, 74, 110)
        elif p["type"] == "GOV":
            card.fill.fore_color.rgb = RGBColor(248, 250, 252)
            card.line.color.rgb = RGBColor(2, 132, 199)
            card.line.width = Pt(1.5)
            title_rgb = RGBColor(15, 23, 42)
        else:
            card.fill.fore_color.rgb = RGBColor(255, 255, 255) # White
            card.line.color.rgb = RGBColor(203, 213, 225) # Slate border
            card.line.width = Pt(1)
            title_rgb = RGBColor(30, 41, 59)
            
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.08)
        ctf.margin_top = Inches(0.06)
        ctf.margin_bottom = Inches(0.04)
        
        # Header line
        cp0 = ctf.paragraphs[0]
        cp0.text = f"{p['id']} — {p['name']}"
        cp0.font.size = Pt(8.5)
        cp0.font.bold = True
        cp0.font.color.rgb = title_rgb
        
        # Status line
        cp_stat = ctf.add_paragraph()
        cp_stat.text = f"Status: {p['status']}"
        cp_stat.font.size = Pt(7.0)
        cp_stat.font.bold = True
        if p["type"] == "COMPLETE":
            cp_stat.font.color.rgb = RGBColor(22, 163, 74)
        elif p["type"] == "NEXT":
            cp_stat.font.color.rgb = RGBColor(2, 132, 199)
        else:
            cp_stat.font.color.rgb = RGBColor(100, 116, 139)
            
        # Bullets
        for b in p["bullets"]:
            cp_b = ctf.add_paragraph()
            cp_b.text = b
            cp_b.font.size = Pt(6.8)
            cp_b.font.color.rgb = RGBColor(51, 65, 85)
            
        # Result Bottom Line
        cp_res = ctf.add_paragraph()
        cp_res.text = p["res"]
        cp_res.font.size = Pt(6.8)
        cp_res.font.bold = True
        if p["type"] == "COMPLETE":
            cp_res.font.color.rgb = RGBColor(21, 128, 61)
        elif p["type"] == "NEXT":
            cp_res.font.color.rgb = RGBColor(3, 105, 161)
        else:
            cp_res.font.color.rgb = RGBColor(71, 85, 105)

    # -------------------------------------------------------------------------
    # Footer: Objective Mapping (Slide 3 Alignment)
    # -------------------------------------------------------------------------
    foot_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.50), Inches(5.95), Inches(12.33), Inches(1.20))
    foot_box.fill.solid()
    foot_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    foot_box.line.color.rgb = RGBColor(148, 163, 184)
    foot_box.line.width = Pt(1)
    
    ftf = foot_box.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = Inches(0.08)
    
    fp0 = ftf.paragraphs[0]
    fp0.text = "ORIGINAL 8 RESEARCH OBJECTIVES MAPPING (ALIGNMENT TO SLIDE 3)"
    fp0.font.size = Pt(8.5)
    fp0.font.bold = True
    fp0.font.color.rgb = RGBColor(15, 23, 42)
    
    # 8 Objectives Text Lines
    obj_lines = [
        "Obj 1 (MuJoCo 6-DoF Integration) → P0 COMPLETE ✅   |   Obj 2 (LSM + 4-Factor Endocrine) → P1 (Next / Progressive)   |   Obj 3 (STDP Continual Plasticity) → P1   |   Obj 4 (Hippocampal Grid Cells) → P2",
        "Obj 5 (Active Inference / Free Energy) → P0 COMPLETE ✅   |   Obj 6 (Physarum Routing) → P3   |   Obj 7 (Mirror-Neuron / Social) → P4   |   Obj 8 (30Hz Validation) → P0 Baseline ✅ (P1+ Ongoing)"
    ]
    for ol in obj_lines:
        fp_ol = ftf.add_paragraph()
        fp_ol.text = ol
        fp_ol.font.size = Pt(7.2)
        fp_ol.font.bold = True
        fp_ol.font.color.rgb = RGBColor(51, 65, 85)
        
    out_pptx_path = os.path.join(REPO_ROOT, "CARL_SLIDE_19_DEVELOPMENT_ROADMAP.pptx")
    prs.save(out_pptx_path)
    print(f"  [OK] Saved Editable PPTX -> {out_pptx_path}")
    return out_pptx_path

# -----------------------------------------------------------------------------
# 3. GENERATE EXPLANATORY ROADMAP NOTES
# -----------------------------------------------------------------------------
def generate_roadmap_notes():
    print(">>> Generating CARL_SLIDE_19_ROADMAP_NOTES.md...")
    notes_content = """# CARL Slide 19: Developmental Phase Roadmap Reference Notes

---

## 1. Core Purpose of Slide 19
Slide 19 bridges the original holistic scientific vision of CARL (presented in Slides 1–18) with the rigorous, phased engineering implementation roadmap required for scientific validation.

### Key Pedagogical Points
1. **The Full Vision is Long-Term**: The comprehensive biological architecture (LSM reservoir, STDP learning, hippocampal grid cells, Physarum route optimization, social mirror neurons, collective swarms) represents the complete multi-year vision of CARL.
2. **Implementation is Progressive**: Rather than claiming all biological mechanisms were instantiated simultaneously from day one, CARL follows a strict, falsifiable phased developmental progression ($P0 \to P1 \to \dots \to P8$).
3. **P0 Foundation is Validated & Complete**: The 6-DoF MuJoCo physical embodiment, Active Inference decision engine ($G(a)$ minimization over 14 candidate arcs), interoceptive homeostasis, HDC episodic memory, and operational safety veto are fully certified (81/81 regression tests PASS, 0 collisions, 0 gate violations).
4. **P1 is the Immediate Next Phase**: Phase P1 focuses on *Individual Development & Emergence*—introducing experience-driven learning via prediction errors, STDP synaptic plasticity, LSM liquid state reservoir dynamics, and 4-factor endocrine modulation under a controlled emergence laboratory.
5. **No False Claims of Premature Integration**: Future biological modules ($P2$ to $P8$) are transparently labeled as **ROADMAP / FUTURE**, maintaining strict scientific integrity before academic judges and seminar audiences.

---

## 2. Comprehensive Phase Definitions

| Phase | Phase Name | Status | Key Focus Areas & Mechanisms | Expected Outcome |
| :--- | :--- | :--- | :--- | :--- |
| **P0** | **Embodied Foundation** | **COMPLETE ✅** | 6-DoF MuJoCo physics, 15-ray sonar array, Active Inference ($G(a)$ rollout evaluation), homeostatic core, online friction adaptation, safety supervisor, fail-closed regression qualification. | Certified stable, non-colliding embodied baseline. |
| **P1** | **Individual Emergence** | **NEXT / UNLOCKED 🔓** | Experience $\to$ prediction error $\to$ adaptation, STDP synaptic plasticity, Liquid State Machine (LSM) spiking reservoir, 4-factor endocrine kinetics (DA, 5-HT, NE, Cortisol). | Experience-dependent behavioral development without hardcoding. |
| **P2** | **Sensory & Spatial** | **FUTURE / ROADMAP** | Latent visual world models, active optical gaze scanning, spatial representation, hippocampal place-cell and grid-cell firing lattices. | Latent visual perception and internal spatial mapping. |
| **P3** | **Adaptive Routing** | **FUTURE / ROADMAP** | Continuous trajectory optimization, long-horizon path planning, multi-step counterfactual rollouts, Physarum-inspired plasmodial tubular route reinforcement. | Optimal pathfinding in complex, dynamic mazes. |
| **P4** | **Social CARL** | **FUTURE / ROADMAP** | Multi-agent co-existence arenas, peer observation, non-verbal acoustic/visual signaling, mirror-neuron-inspired imitation learning. | Inter-organism social learning and peer imitation. |
| **P5** | **Collective Swarm** | **FUTURE / ROADMAP** | Selective shared HDC memory synchronization, swarm hazard broadcasts, collective route discovery, emergent division of labor. | Hive intelligence and distributed multi-agent cooperation. |
| **P6** | **Metacognition** | **FUTURE / ROADMAP** | Failure-pattern self-monitoring, higher-order counterfactual audits, causal attribution, internal epistemic confidence estimation. | Causal self-reflection and introspective error correction. |
| **P7** | **Morphological Co-Evolution** | **FUTURE / ROADMAP** | Brain-directed morphological adaptation, body-brain developmental co-tuning, evolutionary sensor/actuator placement. | Embodied co-evolution adapted to diverse environmental niches. |
| **P8** | **Physical Embodiment** | **FUTURE / ROADMAP** | Sim-to-real parameter transfer, physical brushless DC motor hardware, embedded STM32/Jetson deployment, real-world physical companion validation. | Physical hardware realization in the real world. |

---

## 3. Mapping of the Original 8 Slide 3 Research Objectives

| Original Slide 3 Objective | Assigned Phase | Status & Interpretation |
| :--- | :--- | :--- |
| **Objective 1: Integrated CARL / MuJoCo 6-DoF** | **Phase P0** | **COMPLETE ✅** (Fully embodied in MuJoCo 3.3.4 physics). |
| **Objective 2: LSM Reservoir + 4-Factor Endocrine** | **Phase P1** | **NEXT / UNLOCKED** (Progressively integrated and evaluated in P1). |
| **Objective 3: STDP Continual Synaptic Plasticity** | **Phase P1** | **NEXT / UNLOCKED** (Implemented and evaluated via causal ablation in P1). |
| **Objective 4: Hippocampal Grid & Place Cells** | **Phase P2** | **FUTURE ROADMAP** (Formal spatial development phase). |
| **Objective 5: Free-Energy Active Inference** | **Phase P0** | **COMPLETE ✅** (30 Hz generative self-model minimizing $G(a)$). |
| **Objective 6: Physarum Route Reinforcement** | **Phase P3** | **FUTURE ROADMAP** (Advanced routing and tubular network flow). |
| **Objective 7: Mirror-Neuron / Social Imitation** | **Phase P4** | **FUTURE ROADMAP** (Multi-organism social learning phase). |
| **Objective 8: 30 Hz Digital-Twin Validation** | **P0 Foundation ✅ / P1+ Ongoing** | **P0 BASELINE COMPLETE ✅** (30 Hz loop certified; emergence validation continues in P1+). |

---

## 4. Visual Status Legend
- **GREEN (`#16a34a`)**: **IMPLEMENTED / VALIDATED** — Applied exclusively to Phase P0.
- **BLUE (`#0284c7`)**: **NEXT PHASE / UNLOCKED** — Applied exclusively to Phase P1.
- **SLATE GRAY (`#64748b` / `#cbd5e1`)**: **ROADMAP / FUTURE** — Applied to Phases P2 through P8.
"""
    notes_path = os.path.join(REPO_ROOT, "CARL_SLIDE_19_ROADMAP_NOTES.md")
    with open(notes_path, "w", encoding="utf-8") as f:
        f.write(notes_content)
    print(f"  [OK] Saved Roadmap Reference Notes -> {notes_path}")
    return notes_path

def main():
    print("=================================================================")
    print("   CARL-DX SLIDE 19 DEVELOPMENTAL ROADMAP GENERATOR")
    print("=================================================================")
    png_file = generate_roadmap_png()
    pptx_file = generate_roadmap_pptx()
    notes_file = generate_roadmap_notes()
    
    # Verification
    print("\n>>> Verifying Generated Deliverables...")
    assert os.path.exists(png_file), f"PNG missing: {png_file}"
    assert os.path.exists(pptx_file), f"PPTX missing: {pptx_file}"
    assert os.path.exists(notes_file), f"Notes missing: {notes_file}"
    
    png_size = os.path.getsize(png_file) / (1024 * 1024)
    pptx_size = os.path.getsize(pptx_file) / 1024
    
    print(f"  [VERIFIED] PNG Size:  {png_size:.2f} MB")
    print(f"  [VERIFIED] PPTX Size: {pptx_size:.2f} KB")
    print("\nAll deliverables generated and verified successfully!")

if __name__ == "__main__":
    main()
