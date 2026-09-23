import pptx
from pptx.util import Inches, Pt
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = pptx.Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
t = s.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table

def set_cell_border(cell, color="000000", width="12700"):
    tcPr = cell._tc.get_or_add_tcPr()
    for border_name in ['lnL', 'lnR', 'lnT', 'lnB']:
        ln = parse_xml(f'<a:{border_name} {nsdecls("a")} w="{width}" cmpd="s"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:{border_name}>')
        tcPr.append(ln)

set_cell_border(t.cell(0, 0))
prs.save(r'D:\ebca\test_table.pptx')
print('Table border helper succeeded!')
