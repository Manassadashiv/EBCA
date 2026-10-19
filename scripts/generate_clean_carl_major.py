"""
generate_clean_carl_major.py
Full automated master thesis generator for CARL-DX, strictly benchmarked against:
C:\Users\MANAS\Downloads\project_report_2.pdf
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOC_PATH = r"D:\carl major.docx"
MEDIA_DIR = r"D:\ebca\scratch\docx_media"

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def add_clean_table(doc, headers, data, col_widths=None, alignment=WD_TABLE_ALIGNMENT.CENTER):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = alignment
    hdr_row = tbl.rows[0]
    hdr_tr = hdr_row._tr.get_or_add_trPr()
    hdr_tr.append(OxmlElement('w:tblHeader'))

    for idx, h in enumerate(headers):
        c = hdr_row.cells[idx]
        c.text = h
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.bold = True
        set_cell_shading(c, 'E0E0E0')
        set_cell_margins(c, top=70, bottom=70, left=90, right=90)

    for r_idx, row_vals in enumerate(data):
        row = tbl.rows[r_idx + 1]
        for c_idx, val in enumerate(row_vals):
            c = row.cells[c_idx]
            c.text = str(val)
            p = c.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(9)
            if c_idx == 0:
                p.runs[0].font.bold = True
            set_cell_margins(c, top=50, bottom=50, left=80, right=80)

    if col_widths:
        for r in tbl.rows:
            for idx, w in enumerate(col_widths):
                r.cells[idx].width = w

    return tbl

def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run()
    r.add_break(docx.enum.text.WD_BREAK.PAGE)
    return p

def add_chapter_title(doc, ch_num, ch_title):
    add_page_break(doc)
    p = doc.add_paragraph()
    p.style = doc.styles['Heading 1']
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.keep_with_next = True
    
    r1 = p.add_run(f"CHAPTER-{ch_num}\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(14)
    r1.font.bold = True
    
    r2 = p.add_run(ch_title.upper())
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(14)
    r2.font.bold = True
    return p

def add_heading_2(doc, sec_num, sec_title):
    p = doc.add_paragraph()
    p.style = doc.styles['Heading 2']
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{sec_num}  {sec_title.upper()}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    return p

def add_heading_3(doc, subsec_num, subsec_title):
    p = doc.add_paragraph()
    p.style = doc.styles['Heading 3']
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{subsec_num}  {subsec_title}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.bold = True
    return p

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.style = doc.styles['Normal']
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    return p

def add_equation_p(doc, eq_text, eq_num):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{eq_text}    ({eq_num})")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.italic = True
    return p

def add_image_figure(doc, img_name, caption_text, width=Inches(5.6)):
    img_path = os.path.join(MEDIA_DIR, img_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        r = p_img.add_run()
        r.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        r_cap.font.bold = True
        return p_img, p_cap
    else:
        print(f"Warning: Image {img_name} not found in {MEDIA_DIR}.")
        return None, None

def add_table_caption(doc, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(caption_text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.font.bold = True
    return p

print("Loaded generator definitions.")
